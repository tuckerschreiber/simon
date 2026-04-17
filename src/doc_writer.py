"""Render an AccountStrategy to a .docx file."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from strategy_generator import AccountStrategy

PANW_ORANGE = RGBColor(0xFA, 0x58, 0x2C)


def _slugify(name: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", name.strip().lower()).strip("_")
    return slug or "prospect"


def write_account_doc(
    org: dict[str, Any],
    contacts: list[dict[str, Any]],
    competitors: list[str],
    strategy: AccountStrategy,
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{_slugify(org.get('name', 'prospect'))}_account_plan.docx"

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(f"Account Strategy Plan — {org.get('name', 'Prospect')}", level=0)
    for run in title.runs:
        run.font.color.rgb = PANW_ORANGE

    doc.add_paragraph("Palo Alto Networks | Prospect Outreach Plan").italic = True

    doc.add_heading("1. Company Overview", level=1)
    facts = [
        ("Website", org.get("website_url") or org.get("primary_domain") or "—"),
        ("Industry", org.get("industry") or "—"),
        ("Employees", str(org.get("estimated_num_employees") or "—")),
        (
            "Location",
            ", ".join(
                p for p in [org.get("city"), org.get("state"), org.get("country")] if p
            )
            or "—",
        ),
        ("LinkedIn", org.get("linkedin_url") or "—"),
    ]
    table = doc.add_table(rows=len(facts), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(facts):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v

    doc.add_paragraph(strategy.account_summary)

    doc.add_heading("2. Competitive Landscape & Displacement Angle", level=1)
    if competitors:
        p = doc.add_paragraph()
        p.add_run("Competitor tech detected: ").bold = True
        p.add_run(", ".join(competitors))
    doc.add_paragraph(strategy.competitive_landscape)

    doc.add_heading("3. Pain Hypotheses", level=1)
    for pain in strategy.pain_hypotheses:
        doc.add_paragraph(pain, style="List Bullet")

    doc.add_heading("4. PANW Product Fit", level=1)
    for fit in strategy.panw_product_fit:
        doc.add_paragraph(fit, style="List Bullet")

    doc.add_heading("5. Recommended Entry Point", level=1)
    doc.add_paragraph(strategy.recommended_entry_point)

    doc.add_heading("6. Strategic Rationale (Why Now)", level=1)
    doc.add_paragraph(strategy.strategic_rationale)

    doc.add_heading("7. Target Contacts", level=1)
    if contacts:
        ctable = doc.add_table(rows=1, cols=4)
        ctable.style = "Light Grid Accent 1"
        hdr = ctable.rows[0].cells
        hdr[0].text = "Name"
        hdr[1].text = "Title"
        hdr[2].text = "Email"
        hdr[3].text = "LinkedIn"
        for c in contacts:
            row = ctable.add_row().cells
            row[0].text = f"{c.get('first_name', '')} {c.get('last_name', '')}".strip()
            row[1].text = c.get("title") or ""
            row[2].text = c.get("email") or "—"
            row[3].text = c.get("linkedin_url") or "—"
    else:
        doc.add_paragraph("No contacts identified in this pass.")

    doc.add_heading("8. Draft Outreach Emails", level=1)
    for email in strategy.draft_emails:
        header = doc.add_paragraph()
        header.add_run(f"To: {email.to_persona}").bold = True
        if email.to_name:
            header.add_run(f"  ({email.to_name})")
        subj = doc.add_paragraph()
        subj.add_run("Subject: ").bold = True
        subj.add_run(email.subject)
        doc.add_paragraph(email.body)
        rationale = doc.add_paragraph()
        rationale.add_run("Why this angle: ").italic = True
        rationale.add_run(email.rationale).italic = True
        doc.add_paragraph()

    doc.add_heading("9. Next Steps", level=1)
    for step in strategy.next_steps:
        doc.add_paragraph(step, style="List Number")

    doc.save(path)
    return path
